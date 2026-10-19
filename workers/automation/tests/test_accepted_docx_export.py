"""Owned native-schema exports read and edited with an independent DOCX library."""

from __future__ import annotations

import hashlib
import io
import json
import os
import sqlite3
import subprocess
import sys
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.dom.minidom import parseString
from xml.etree import ElementTree as ET

import pytest
from docx import Document

from jobctrl.domain.materials.editable_export import ExportFailure
from jobctrl.domain.materials.services import ResumeAssembler
from jobctrl.infrastructure.materials import accepted_export as export
from jobctrl.infrastructure.materials.docx_export import CP, VT
from jobctrl.infrastructure.migrations.schema_v14 import create_exact_v14_schema
from tests.test_resume_assembler import _payload, _profile

JOB = "90000000-0000-4000-8000-000000001030"


def synthetic_resume():
    profile = _profile()
    profile["personal"].update(full_name="Zoë Example", email="zoe@example.test", phone="+34 555 010 010")
    profile["resume"]["experience_entries"][0].update(date_range="févr. 2020 – présent", company="Équipe & Co")
    payload = _payload()
    payload["experience_updates"][0]["bullets"] = ["Delivered 25% improvement;  €1.2M retained.", "Second accepted fact."]
    return ResumeAssembler().assemble_resume_text(payload, profile)


@pytest.fixture
def owned(tmp_path):
    root = tmp_path / "owned-workspace"
    root.mkdir()
    connection = sqlite3.connect(root / "jobctrl.db")
    create_exact_v14_schema(connection)
    connection.execute("INSERT INTO jobs(tenant_id,job_id,url,title,company,apply_status,applied_at) "
                       "VALUES('local',?,'https://example.test/job','Owned job','Synthetic Employer','applied','2026-10-01')", (JOB,))
    connection.execute("INSERT INTO candidate_profiles(tenant_id,profile_id,version,personal_full_name,updated_at) "
                       "VALUES('local','default',9,'Different current profile','2026-10-17')")
    connection.execute("INSERT INTO job_materials(tenant_id,job_id,generation,status,created_at,updated_at) "
                       "VALUES('local',?,1,'approved','2026-10-01','2026-10-01')", (JOB,))
    connection.execute("INSERT INTO application_review_decisions(tenant_id,decision_id,job_id,decision,decided_at,materials_generation) "
                       "VALUES('local','owned-approval',?,'approved','2026-10-01',1)", (JOB,))
    texts = {"tailored_resume": synthetic_resume(),
             "cover_letter": "  Madame,\r\n\r\nJ’ai livré 25% en 2020.\t€1.2M & <literal>.  \r\n\r\nZoë Example\r\n"}
    for kind, text in texts.items():
        path = root / f"{kind}.txt"
        path.write_bytes(text.encode())
        metadata = {"locale": "fr-FR", "claim_verification_id": "verification-owned", "line_anchors": [
            {"line_id": "experience:acme_swe#0", "evidence_ids": ["evidence-owned"]}], "private_prompt": "never embed me"}
        connection.execute("INSERT INTO job_materials_artifacts(tenant_id,job_id,generation,artifact_type,artifact_id,"
                           "status,path,render_format,created_at,metadata_json) VALUES('local',?,1,?,?,'approved',?,'text',?,?)",
                           (JOB, kind, kind, str(path), "2026-10-01", json.dumps(metadata)))
    connection.commit()
    yield connection, root
    connection.close()


def run(owned, kind="tailored_resume", **overrides):
    connection, root = owned
    args = dict(tenant_id="local", job_id=JOB, artifact_id=kind, generation=1, root=root)
    args.update(overrides)
    return export.export_accepted_material(connection, **args)


@pytest.mark.parametrize("kind", ["tailored_resume", "cover_letter"])
def test_carried_source_exports_each_exact_generation_without_ambiguity(owned, kind):
    connection, _ = owned
    for generation in (2, 3):
        connection.execute("INSERT INTO job_materials(tenant_id,job_id,generation,status,created_at,updated_at) "
                           "VALUES('local',?,?,'approved','now','now')", (JOB, generation))
        connection.execute("INSERT INTO job_materials_artifacts "
                           "SELECT tenant_id,job_id,?,artifact_type,artifact_id,status,path,render_format,"
                           "size_bytes,metadata_json,created_at,superseded_at FROM job_materials_artifacts "
                           "WHERE tenant_id='local' AND job_id=? AND generation=1", (generation, JOB))
    connection.commit()
    before = preservation_state(owned)
    for generation in (1, 2, 3):
        receipt = run(owned, kind, generation=generation)
        assert receipt["generation"] == generation
        assert Document(receipt["path"]).core_properties.identifier == kind
    assert preservation_state(owned) == before


def preservation_state(owned):
    connection, root = owned
    return {
        "sources": [tuple(row) for row in connection.execute("SELECT * FROM job_materials_artifacts")],
        "materials": [tuple(row) for row in connection.execute("SELECT * FROM job_materials")],
        "jobs": [tuple(row) for row in connection.execute("SELECT * FROM jobs")],
        "profile": [tuple(row) for row in connection.execute("SELECT * FROM candidate_profiles")],
        "reviews": [tuple(row) for row in connection.execute("SELECT * FROM application_review_decisions")],
        "files": {str(path): path.read_bytes() for path in root.glob("*.txt")},
    }


@pytest.mark.parametrize("kind", ["tailored_resume", "cover_letter"])
@pytest.mark.parametrize("paper", ["A4", "Letter"])
def test_native_export_round_trips_every_line_and_is_editable(owned, kind, paper):
    before = preservation_state(owned)
    receipt = run(owned, kind, settings={"paper": paper})
    path = Path(receipt["path"])
    # python-docx uses an independent OPC parser and lxml, not our XML writer.
    document = Document(path)
    expected = (owned[1] / f"{kind}.txt").read_bytes().decode().replace("\r\n", "\n").split("\n")
    assert [p.text for p in document.paragraphs] == expected
    assert document.core_properties.language == "fr-FR"
    assert document.core_properties.identifier == kind
    assert document.core_properties.author == ""
    assert round(document.sections[0].page_width.inches, 2) == (8.27 if paper == "A4" else 8.5)
    if kind == "tailored_resume":
        assert [p.text for p in document.paragraphs if p.style.name == "Heading 1"] == [
            "EXECUTIVE PROFILE", "EXPERIENCE", "EDUCATION", "SKILLS"]
        assert document.paragraphs[0].style.name == "Title"
        assert len([p for p in document.paragraphs if p.style.name == "List Paragraph"]) == 2
    assert receipt["sourceSha256"] == hashlib.sha256((owned[1] / f"{kind}.txt").read_bytes()).hexdigest()
    assert receipt["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert path.stat().st_mode & 0o777 == 0o600
    with zipfile.ZipFile(path) as archive:
        # LibreOffice's package detector refused these OPC roots when prefixed;
        # independent text readers alone did not expose that interoperability bug.
        for part in ("[Content_Types].xml", "_rels/.rels", "word/_rels/document.xml.rels"):
            assert parseString(archive.read(part)).documentElement.prefix is None
        custom = ET.fromstring(archive.read("docProps/custom.xml"))
        binding = json.loads(custom.find(f"{{{CP}}}property/{{{VT}}}lpwstr").text)
        assert binding["sourceArtifactId"] == kind and binding["generation"] == 1
        assert binding["sourceMetadataSha256"] == receipt["sourceMetadataSha256"]
        assert binding["settings"] == receipt["settings"]
        assert all(b"never embed me" not in archive.read(name) for name in archive.namelist())
    # Prove editable text, rather than pictures or an embedded PDF.
    document.paragraphs[-1].add_run(" Owned local edit")
    saved = io.BytesIO()
    document.save(saved)
    assert [p.text for p in Document(saved).paragraphs] == expected[:-1] + [expected[-1] + " Owned local edit"]
    assert preservation_state(owned) == before
    assert owned[0].execute("SELECT artifact_type,status FROM job_artifacts").fetchall() == [("material_docx", "approved")]


def test_source_identity_settings_and_bytes_are_bound_without_refreshing_profile(owned):
    first = run(owned)
    second = run(owned)
    other = run(owned, settings={"paper": "Letter"})
    assert first["path"] != second["path"]
    assert first["sha256"] == second["sha256"]
    assert first["sha256"] != other["sha256"]
    assert first["settingsSha256"] != other["settingsSha256"]
    assert "Different current profile" not in "\n".join(p.text for p in Document(first["path"]).paragraphs)


@pytest.mark.parametrize(("settings", "code"), [
    ({"format": "pdf"}, "unsupported_export_format"),
    ({"layout": "pixel-perfect"}, "unsupported_export_settings"),
    ({"paper": "Legal"}, "unsupported_export_settings"),
    ({"font": "custom"}, "unsupported_export_settings"),
    ({"locale": "en-US"}, "source_locale_mismatch"),
    ({"locale": "bad_locale"}, "unsupported_export_settings"),
])
def test_unsupported_formatting_and_locale_are_explicit_and_read_only(owned, settings, code):
    before = preservation_state(owned)
    with pytest.raises(ExportFailure, match=code):
        run(owned, settings=settings)
    assert preservation_state(owned) == before
    assert not list(owned[1].rglob("*.docx"))


def test_unknown_structure_is_literal_and_unknown_locale_is_not_guessed(owned):
    text = "  CV Personnel  \nRésumé\nDate 01/02/2020\n- Élément précis\n"
    (owned[1] / "tailored_resume.txt").write_text(text)
    owned[0].execute("UPDATE job_materials_artifacts SET metadata_json='{}'")
    owned[0].commit()
    result = run(owned)
    doc = Document(result["path"])
    assert [p.text for p in doc.paragraphs] == text.split("\n")
    assert all(p.style.name == "Normal" for p in doc.paragraphs)
    assert doc.core_properties.language == "und"
    assert any("Unrecognized resume structure" in warning for warning in result["warnings"])
    assert any("undetermined" in warning for warning in result["warnings"])
    assert Document(run(owned, settings={"locale": "ca"})["path"]).core_properties.language == "ca"


@pytest.mark.parametrize(("change", "code"), [
    ("candidate", "accepted_source_unavailable"), ("superseded", "accepted_source_unavailable"),
    ("foreign_tenant", "accepted_source_unavailable"), ("foreign_job", "accepted_source_unavailable"),
    ("generation", "stale_source_generation"), ("pdf", "unsupported_source_format"),
    ("control", "unsupported_source_characters"), ("large", "source_too_large"),
])
def test_unaccepted_stale_and_unsupported_sources_fail_closed(owned, change, code):
    kwargs = {}
    if change in {"candidate", "superseded"}:
        owned[0].execute("UPDATE job_materials_artifacts SET status=?", (change,))
    elif change == "foreign_tenant":
        kwargs["tenant_id"] = "other"
    elif change == "foreign_job":
        kwargs["job_id"] = "90000000-0000-4000-8000-000000001031"
    elif change == "generation":
        kwargs["generation"] = 2
    elif change == "pdf":
        owned[0].execute("UPDATE job_materials_artifacts SET render_format='html_pdf'")
    else:
        (owned[1] / "tailored_resume.txt").write_text("x\x00y" if change == "control" else "x" * 200001)
    owned[0].commit()
    before = preservation_state(owned)
    with pytest.raises(ExportFailure, match=code):
        run(owned, **kwargs)
    assert preservation_state(owned) == before
    assert not list(owned[1].rglob("*.docx"))


@pytest.mark.parametrize("change", ["bytes", "metadata", "status", "path"])
def test_concurrent_source_change_during_render_is_fenced(owned, change):
    original = export.render_docx
    def render(*args):
        result = original(*args)
        if change == "bytes":
            (owned[1] / "tailored_resume.txt").write_text("Successor-owned material")
        else:
            connection = sqlite3.connect(owned[1] / "jobctrl.db")
            if change == "metadata":
                connection.execute("UPDATE job_materials_artifacts SET metadata_json='{}'")
            elif change == "status":
                connection.execute("UPDATE job_materials_artifacts SET status='superseded'")
            else:
                connection.execute("UPDATE job_materials_artifacts SET path=?", (str(owned[1] / "cover_letter.txt"),))
            connection.commit()
            connection.close()
        return result
    with patch.object(export, "render_docx", side_effect=render), pytest.raises(ExportFailure):
        run(owned)
    assert not list(owned[1].rglob("*.docx"))
    assert owned[0].execute("SELECT count(*) FROM job_artifacts").fetchone()[0] == 0


@pytest.mark.parametrize("failure", ["renderer", "corrupt", "write", "registry"])
def test_failed_export_preserves_previous_files_approval_and_application_state(owned, failure):
    previous = run(owned)
    previous_bytes = Path(previous["path"]).read_bytes()
    before = preservation_state(owned)
    if failure == "registry":
        owned[0].execute("CREATE TRIGGER owned_failure BEFORE INSERT ON job_artifacts BEGIN SELECT RAISE(ABORT,'owned'); END")
        owned[0].commit()
        with pytest.raises(sqlite3.IntegrityError):
            run(owned)
    else:
        target = "_write_exclusive" if failure == "write" else "render_docx"
        kwargs = {"return_value": b"not a DOCX"} if failure == "corrupt" else {"side_effect": OSError("owned")}
        with patch.object(export, target, **kwargs), pytest.raises((OSError, ExportFailure)):
            run(owned)
    assert preservation_state(owned) == before
    assert Path(previous["path"]).read_bytes() == previous_bytes
    assert list(owned[1].rglob("*.docx")) == [Path(previous["path"])]
    assert owned[0].execute("SELECT count(*) FROM job_artifacts").fetchone()[0] == 1


@pytest.mark.parametrize("kind", ["source", "output"])
def test_symlinks_cannot_escape_owned_workspace(owned, tmp_path, kind):
    outside = tmp_path / "outside"
    outside.mkdir()
    if kind == "source":
        source = owned[1] / "tailored_resume.txt"
        source.unlink()
        target = outside / "private.txt"
        target.write_text("Never read this")
        source.symlink_to(target)
    else:
        (owned[1] / "generated").symlink_to(outside, target_is_directory=True)
    with pytest.raises(ExportFailure, match="artifact_path_outside_workspace"):
        run(owned)
    assert not list(owned[1].rglob("*.docx"))
    assert not list(outside.glob("*.docx"))


def test_caller_transaction_is_never_committed_or_rolled_back(owned):
    owned[0].execute("UPDATE jobs SET title='Caller-owned change'")
    with pytest.raises(ExportFailure, match="export_requires_own_transaction"):
        run(owned)
    assert owned[0].in_transaction
    assert owned[0].execute("SELECT title FROM jobs").fetchone()[0] == "Caller-owned change"
    owned[0].rollback()


def test_source_file_change_after_staging_removes_only_new_export(owned):
    prior = run(owned)
    prior_bytes = Path(prior["path"]).read_bytes()
    write = export._write_exclusive
    def change(path, data):
        write(path, data)
        (owned[1] / "tailored_resume.txt").write_text("Concurrent source owner")
    with patch.object(export, "_write_exclusive", side_effect=change), pytest.raises(ExportFailure, match="stale_export_source"):
        run(owned)
    assert list(owned[1].rglob("*.docx")) == [Path(prior["path"])]
    assert Path(prior["path"]).read_bytes() == prior_bytes
    assert (owned[1] / "tailored_resume.txt").read_text() == "Concurrent source owner"
    assert owned[0].execute("SELECT count(*) FROM job_artifacts").fetchone()[0] == 1


def test_exclusive_collision_never_removes_a_preexisting_file(owned):
    prior = run(owned)
    prior_bytes = Path(prior["path"]).read_bytes()
    with patch.object(export.uuid, "uuid4") as identifier:
        identifier.return_value.hex = Path(prior["path"]).stem
        with pytest.raises(FileExistsError):
            run(owned)
    assert Path(prior["path"]).read_bytes() == prior_bytes
    assert owned[0].execute("SELECT count(*) FROM job_artifacts").fetchone()[0] == 1


def test_actual_cli_subprocess_exports_and_reports_unsupported_format(owned):
    command = [sys.executable, "-m", "jobctrl.cli", "material-export", JOB,
               "--source-artifact-id", "cover_letter", "--generation", "1"]
    env = {**os.environ, "JOBCTRL_DIR": str(owned[1])}
    before = preservation_state(owned)
    completed = subprocess.run(command, env=env, text=True, capture_output=True, timeout=30)
    assert completed.returncode == 0, completed.stderr
    result = json.loads(completed.stdout)
    assert Document(result["path"]).core_properties.language == "fr-FR"
    refused = subprocess.run(command + ["--format", "rtf"], env=env, text=True, capture_output=True, timeout=30)
    assert refused.returncode == 1 and "unsupported_export_format" in refused.stderr
    assert preservation_state(owned) == before
